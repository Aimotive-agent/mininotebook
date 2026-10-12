"""MiniNotebook — a local NotebookLM-style app (sources, grounded chat, Studio artifacts)."""
import base64
import io
import json
import time
import os
import sqlite3
import uuid
from typing import Optional

import markdown
import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from ebooklib import epub
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pptx import Presentation
from pydantic import BaseModel
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text
from youtube_transcript_api import YouTubeTranscriptApi

import audio
import rag
import video
from openai import OpenAI

DATA_DIR = os.getenv("DATA_DIR", "/data")
DB = os.path.join(DATA_DIR, "mininotebook.db")
FILES_DIR = os.path.join(DATA_DIR, "files")
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(FILES_DIR, exist_ok=True)

app = FastAPI(title="MiniNotebook")


@app.get("/")
def index():
    return FileResponse(os.path.join(os.path.dirname(__file__), "web", "index.html"))


def _db():
    c = sqlite3.connect(DB)
    c.row_factory = sqlite3.Row
    c.execute("CREATE TABLE IF NOT EXISTS notebooks(id TEXT PRIMARY KEY, title TEXT, created TEXT)")
    try:
        c.execute("ALTER TABLE notebooks ADD COLUMN memory TEXT")
    except Exception:
        pass
    c.execute("CREATE TABLE IF NOT EXISTS sources(id TEXT PRIMARY KEY, notebook_id TEXT, name TEXT, kind TEXT, created TEXT)")
    c.execute("CREATE TABLE IF NOT EXISTS notes(id TEXT PRIMARY KEY, notebook_id TEXT, body TEXT, created TEXT)")
    c.execute("CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, notebook_id TEXT, kind TEXT, title TEXT, content TEXT, created TEXT)")
    c.execute("CREATE TABLE IF NOT EXISTS messages(id TEXT PRIMARY KEY, notebook_id TEXT, role TEXT, content TEXT, citations TEXT, created TEXT)")
    try:
        c.execute("ALTER TABLE messages ADD COLUMN verification TEXT")
    except Exception:
        pass
    return c


def _llm():
    return OpenAI(base_url=os.getenv("LLM_BASE", "https://api.deepseek.com/v1"), api_key=os.getenv("LLM_KEY", ""))


def _complete(prompt, max_tokens=2500, temperature=0.35):
    last = "unknown"
    for _ in range(3):
        try:
            r = _llm().chat.completions.create(
                model=os.getenv("LLM_MODEL", "deepseek-v4-pro"),
                messages=[{"role": "user", "content": prompt}], temperature=temperature, max_tokens=max_tokens,
                extra_body={"thinking": {"type": "disabled"}},
            )
            m = r.choices[0].message
            c = (m.content or "").strip()
            if not c:
                rc = getattr(m, "reasoning_content", None)
                if rc is None and hasattr(m, "model_extra"):
                    rc = (m.model_extra or {}).get("reasoning_content")
                c = (rc or "").strip()
            if c:
                return c
            last = "empty response"
        except Exception as e:  # noqa: BLE001
            last = str(e)[:200]
        time.sleep(2)
    raise HTTPException(502, f"Generation failed after 3 attempts ({last}). Try again.")


class NotebookIn(BaseModel):
    title: str


class RenameIn(BaseModel):
    title: str


class ChatIn(BaseModel):
    question: str


class TextIn(BaseModel):
    title: str
    text: str


class NoteIn(BaseModel):
    body: str


class StudioIn(BaseModel):
    kind: str


# ---------------- notebooks ----------------

@app.get("/api/notebooks")
def list_notebooks():
    c = _db()
    out = []
    for r in c.execute("SELECT * FROM notebooks ORDER BY created DESC").fetchall():
        n = c.execute("SELECT COUNT(*) n FROM sources WHERE notebook_id=?", (r["id"],)).fetchone()["n"]
        d = c.execute("SELECT COUNT(*) n FROM documents WHERE notebook_id=?", (r["id"],)).fetchone()["n"]
        out.append({**dict(r), "sources": n, "artifacts": d})
    return out


@app.post("/api/notebooks")
def create_notebook(body: NotebookIn):
    nid = uuid.uuid4().hex
    c = _db()
    c.execute("INSERT INTO notebooks VALUES (?,?,datetime('now'))", (nid, body.title))
    c.commit()
    return {"id": nid, "title": body.title}


@app.patch("/api/notebooks/{nb_id}")
def rename_notebook(nb_id: str, body: RenameIn):
    c = _db()
    c.execute("UPDATE notebooks SET title=? WHERE id=?", (body.title, nb_id))
    c.commit()
    return {"id": nb_id, "title": body.title}


@app.delete("/api/notebooks/{nb_id}")
def delete_notebook(nb_id: str):
    c = _db()
    for s in c.execute("SELECT id FROM sources WHERE notebook_id=?", (nb_id,)).fetchall():
        rag.delete_source(nb_id, s["id"])
    for t in ("sources", "notes", "documents"):
        c.execute(f"DELETE FROM {t} WHERE notebook_id=?", (nb_id,))
    c.execute("DELETE FROM notebooks WHERE id=?", (nb_id,))
    c.commit()
    return {"deleted": nb_id}


# ---------------- sources ----------------

@app.get("/api/notebooks/{nb_id}/sources")
def list_sources(nb_id: str):
    c = _db()
    return [dict(r) for r in c.execute("SELECT * FROM sources WHERE notebook_id=? ORDER BY created", (nb_id,)).fetchall()]


def _vision_ocr(png_b64: str) -> str:
    """OCR one page image via the DeepSeek vision model (deepseek-flash)."""
    r = _llm().chat.completions.create(
        model=os.getenv("OCR_MODEL", "deepseek-flash"),
        messages=[{"role": "user", "content": [
            {"type": "text", "text": "Transcribe ALL text in this document image, preserving line breaks. Output only the text, no commentary."},
            {"type": "image_url", "image_url": {"url": "data:image/png;base64," + png_b64}},
        ]}],
        max_tokens=2500, temperature=0,
        extra_body={"thinking": {"type": "disabled"}},
    )
    m = r.choices[0].message
    return (m.content or m.reasoning_content or "").strip()


def _ocr_pdf(raw: bytes) -> str:
    import fitz  # PyMuPDF
    doc = fitz.open(stream=raw, filetype="pdf")
    pages = []
    for i, page in enumerate(doc):
        if i >= 40:
            break
        pix = page.get_pixmap(dpi=140)
        txt = _vision_ocr(base64.b64encode(pix.tobytes("png")).decode())
        if txt:
            pages.append(txt)
    return "\n\n".join(pages)


def _extract_text(name: str, raw: bytes) -> str:
    low = name.lower()
    if low.endswith(".pdf"):
        text = "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(raw)).pages)
        if text.strip():
            return text
        ocr = _ocr_pdf(raw)
        if ocr.strip():
            return ocr
        raise HTTPException(400, f"'{name}' has no text layer and OCR produced nothing.")
    if low.endswith(".docx"):
        return "\n".join(p.text for p in DocxDocument(io.BytesIO(raw)).paragraphs)
    if low.endswith(".pptx"):
        return "\n".join(sh.text_frame.text for sl in Presentation(io.BytesIO(raw)).slides for sh in sl.shapes if sh.has_text_frame)
    if low.endswith(".epub"):
        return BeautifulSoup("\n".join(i.get_body_content().decode("utf-8", "ignore") for i in epub.read_epub(io.BytesIO(raw)).get_items_of_type(9)), "html.parser").get_text(" ")
    if low.endswith(".rtf"):
        return rtf_to_text(raw.decode("utf-8", "ignore"))
    if low.endswith(".md"):
        return markdown.markdown(raw.decode("utf-8", "ignore"))
    if low.endswith((".txt", ".csv", ".json", ".log")):
        return raw.decode("utf-8", "ignore")
    if low.endswith((".html", ".htm")):
        return BeautifulSoup(raw.decode("utf-8", "ignore"), "html.parser").get_text(" ", strip=True)
    if low.endswith((".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp")):
        import fitz
        try:
            doc = fitz.open(stream=raw, filetype=low.lstrip("."))
            out = []
            for page in doc:
                out.append(_vision_ocr(base64.b64encode(page.get_pixmap(dpi=150).tobytes("png")).decode()))
            txt = "\n".join(x for x in out if x).strip()
            if txt:
                return txt
        except Exception:
            pass
        raise HTTPException(400, f"Couldn't read any text from image '{name}'.")
    raise HTTPException(400, f"Unsupported file type: '{name}'.")


def _fetch_url(url: str) -> str:
    vid = _youtube_id(url)
    if vid:
        try:
            return "\n".join(s.text for s in YouTubeTranscriptApi().fetch(vid))
        except Exception as e:  # noqa: BLE001
            raise HTTPException(400, f"No transcript available for that YouTube video ({e}).") from e
    r = requests.get(url, timeout=30, headers={"User-Agent": "Mozilla/5.0"})
    r.raise_for_status()
    if "pdf" in r.headers.get("Content-Type", ""):
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(r.content)).pages)
    soup = BeautifulSoup(r.text, "html.parser")
    for t in soup(["script", "style", "nav", "footer", "header"]):
        t.decompose()
    txt = soup.get_text(" ", strip=True)
    if txt:
        return txt
    raise HTTPException(400, "No readable text at that URL (it may be JavaScript-only).")


def _youtube_id(url: str):
    import re
    m = re.search(r"(?:v=|youtu\.be/|/shorts/|/live/)([A-Za-z0-9_-]{11})", url)
    return m.group(1) if m else None


def _add(nb_id, name, text, kind):
    sid = uuid.uuid4().hex
    n = rag.index(nb_id, sid, name, text)
    c = _db()
    c.execute("INSERT INTO sources VALUES (?,?,?,?,datetime('now'))", (sid, nb_id, name, kind))
    c.commit()
    return {"id": sid, "name": name, "kind": kind, "chunks": n}


def _store_file(sid, name, raw):
    import pathlib
    ext = pathlib.Path(name).suffix.lower()
    with open(os.path.join(FILES_DIR, sid + ext), "wb") as f:
        f.write(raw)
    return sid + ext


@app.post("/api/notebooks/{nb_id}/sources")
async def add_source(nb_id: str, file: Optional[UploadFile] = File(None), url: Optional[str] = None):
    if file and file.filename:
        raw = await file.read()
        text = _extract_text(file.filename, raw)
        sid = uuid.uuid4().hex
        _store_file(sid, file.filename, raw)
        n = rag.index(nb_id, sid, file.filename, text)
        c = _db()
        c.execute("INSERT INTO sources VALUES (?,?,?,?,datetime('now'))", (sid, nb_id, file.filename, "file"))
        c.commit()
        return {"id": sid, "name": file.filename, "kind": "file", "chunks": n}
    if url:
        return _add(nb_id, url, _fetch_url(url), "url")
    raise HTTPException(400, "Provide a file or a URL.")


@app.post("/api/notebooks/{nb_id}/sources/text")
def add_text(nb_id: str, body: TextIn):
    return _add(nb_id, body.title, body.text, "text")


@app.delete("/api/notebooks/{nb_id}/sources/{sid}")
def remove_source(nb_id: str, sid: str):
    c = _db()
    c.execute("DELETE FROM sources WHERE id=? AND notebook_id=?", (sid, nb_id))
    c.commit()
    rag.delete_source(nb_id, sid)
    return {"deleted": sid}


def _source_or_404(sid):
    r = _db().execute("SELECT * FROM sources WHERE id=?", (sid,)).fetchone()
    if not r:
        raise HTTPException(404, "source not found")
    return dict(r)


@app.get("/api/sources/{sid}/summary")
def source_summary(sid: str):
    s = _source_or_404(sid)
    text = rag.load_text(sid)
    if not text:
        raise HTTPException(404, "no text stored for this source")
    out = _complete(f"Summarize this document in 80-100 words, then list 3-4 key points as bullets.\n\n{text[:7000]}", 2000, 0.2)
    return {"id": sid, "name": s["name"], "summary": out}


@app.get("/api/sources/{sid}/text")
def source_text(sid: str):
    s = _source_or_404(sid)
    return {"id": sid, "name": s["name"], "text": rag.load_text(sid)}


@app.get("/api/sources/{sid}/file")
def source_file(sid: str):
    s = _source_or_404(sid)
    for fn in os.listdir(FILES_DIR):
        if fn.startswith(sid):
            p = os.path.join(FILES_DIR, fn)
            ext = os.path.splitext(fn)[1].lower()
            mt = {
                ".pdf": "application/pdf", ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
                ".gif": "image/gif", ".webp": "image/webp", ".bmp": "image/bmp", ".tif": "image/tiff", ".tiff": "image/tiff",
                ".txt": "text/plain", ".md": "text/plain", ".csv": "text/plain", ".json": "application/json",
                ".html": "text/html", ".htm": "text/html",
            }.get(ext, "application/octet-stream")
            return FileResponse(p, media_type=mt, content_disposition_type="inline", filename=s["name"])
    raise HTTPException(404, "no original file stored (URL/pasted source)")


@app.get("/api/sources/{sid}/related")
def source_related(sid: str, query: str = ""):
    s = _source_or_404(sid)
    if not query.strip():
        return {"matches": []}
    hits = rag.retrieve_by_source(s["notebook_id"], sid, query, top_k=4)
    matches = []
    for h in hits:
        note = _complete(
            f"The user is working on this question in the chat: \"{query}\".\n\n"
            f"This excerpt from a source document is relevant. In one short sentence, explain what in it matches the user's question.\n\nExcerpt: {h['text'][:1200]}",
            800, 0.2)
        matches.append({"excerpt": h["text"][:800], "why": note})
    return {"matches": matches}


# ---------------- chat ----------------

@app.post("/api/notebooks/{nb_id}/chat")
def chat(nb_id: str, body: ChatIn):
    c = _db()
    nb = c.execute("SELECT memory FROM notebooks WHERE id=?", (nb_id,)).fetchone()
    memory = (nb["memory"] or "").strip() if nb else ""
    hits = rag.retrieve(nb_id, body.question, top_k=24)
    if not hits:
        raise HTTPException(404, "No sources in this notebook yet.")
    ctx = "\n\n".join(f"[{i+1}] ({h['name']})\n{h['text'][:1500]}" for i, h in enumerate(hits))
    notes = [r["body"] for r in c.execute("SELECT body FROM notes WHERE notebook_id=? ORDER BY created DESC LIMIT 10", (nb_id,)).fetchall()]
    notes_txt = "\n\n".join(f"- {n[:500]}" for n in notes)
    draft_prompt = (
        "You are assisting a self-represented litigant building a legal case. Answer the question by "
        "deeply analysing ALL the numbered sources. Cite every factual claim with its source number in "
        "brackets like [1][3]. Clearly distinguish (a) established facts, (b) the worker's allegations, "
        "and (c) any formal findings. If sources conflict, point out the conflict. Do not speculate beyond "
        "the sources.\n\n"
        + (f"Your persistent memory for this case (carried from earlier conversations):\n{memory}\n\n" if memory else "")
        + (f"Saved notes (context):\n{notes_txt}\n\n" if notes_txt else "")
        + f"Sources:\n{ctx}\n\nQuestion: {body.question}\n\nAnswer:"
    )
    answer = _complete(draft_prompt, 3000, 0.3)
    fc_prompt = (
        "Fact-check the answer below against the numbered sources. Report: (1) any claim NOT supported by "
        "its cited source, (2) any contradiction between sources, (3) any material fact the answer omitted. "
        "Then write a short 'Verification' note (2-4 sentences) stating whether the answer is fully supported "
        "or flagging the specific issues.\n\n"
        f"Sources:\n{ctx}\n\nAnswer to verify:\n{answer}\n\nVerification:"
    )
    verification = _complete(fc_prompt, 800, 0.2)
    # update persistent memory
    try:
        mem_prompt = (
            "You maintain a persistent memory for a legal case notebook. Given the existing memory and the "
            "latest question + answer, write an updated concise memory (max 300 words) covering: established "
            "facts, key evidence/documents, decisions or conclusions, and open questions. Output only the "
            "memory text, no preamble.\n\n"
            f"Existing memory:\n{memory or '(none)'}\n\nLatest Q: {body.question}\nLatest A: {answer[:1500]}\n\nUpdated memory:"
        )
        new_memory = _complete(mem_prompt, 700, 0.3)
        c.execute("UPDATE notebooks SET memory=? WHERE id=?", (new_memory, nb_id))
    except Exception:
        pass
    citations = [{"number": i+1, "name": h["name"], "excerpt": h["text"][:220]} for i, h in enumerate(hits)]
    c.execute("INSERT INTO messages VALUES (?,?,?,?,?,datetime('now'),?)", (uuid.uuid4().hex, nb_id, "user", body.question, None, None))
    c.execute("INSERT INTO messages VALUES (?,?,?,?,?,datetime('now'),?)", (uuid.uuid4().hex, nb_id, "assistant", answer, json.dumps(citations), verification))
    c.commit()
    return {"answer": answer, "citations": citations, "verification": verification}


@app.get("/api/notebooks/{nb_id}/messages")
def list_messages(nb_id: str):
    c = _db()
    out = []
    for r in c.execute("SELECT role, content, citations, verification, created FROM messages WHERE notebook_id=? ORDER BY created, rowid", (nb_id,)).fetchall():
        out.append({"role": r["role"], "content": r["content"], "citations": json.loads(r["citations"]) if r["citations"] else None, "verification": r["verification"], "created": r["created"]})
    return out


@app.post("/api/notebooks/{nb_id}/suggestions")
def suggestions(nb_id: str):
    srcs = rag.all_source_texts(nb_id)
    if not srcs:
        raise HTTPException(404, "no sources")
    corpus = "\n\n".join(s["text"][:2000] for s in srcs[:4])
    out = _complete("Suggest exactly 3 short questions a user could ask about this material. One per line, no numbering.\n\n" + corpus, 1500, 0.6)
    return {"questions": [q.strip("-•* ") for q in out.splitlines() if q.strip()][:3]}


# ---------------- notes ----------------

@app.get("/api/notebooks/{nb_id}/notes")
def list_notes(nb_id: str):
    return [dict(r) for r in _db().execute("SELECT * FROM notes WHERE notebook_id=? ORDER BY created DESC", (nb_id,)).fetchall()]


@app.post("/api/notebooks/{nb_id}/notes")
def add_note(nb_id: str, body: NoteIn):
    nid = uuid.uuid4().hex
    c = _db()
    c.execute("INSERT INTO notes VALUES (?,?,?,datetime('now'))", (nid, nb_id, body.body))
    c.commit()
    return {"id": nid, "body": body.body}


@app.delete("/api/notebooks/{nb_id}/notes/{nid}")
def del_note(nb_id: str, nid: str):
    c = _db()
    c.execute("DELETE FROM notes WHERE id=? AND notebook_id=?", (nid, nb_id))
    c.commit()
    return {"deleted": nid}


# ---------------- studio (generated artifacts) ----------------

STUDIO = {
    "audio":     ("Audio Overview", "podcasts"),
    "slides":    ("Slide Deck", "slideshow"),
    "video":     ("Video Overview", "movie"),
    "mindmap":   ("Mind Map", "account_tree"),
    "report":    ("Report", "description"),
    "flashcards":("Flashcards", "style"),
    "quiz":      ("Quiz", "quiz"),
    "infographic":("Infographic", "insights"),
    "datatable": ("Data Table", "table_chart"),
    "summary":   ("Summary", "summarize"),
    "study_guide":("Study Guide", "school"),
}

_PROMPTS = {
 "summary": "Write a comprehensive multi-part summary of this workers' compensation case covering: the parties and claim numbers, the injuries and key medical opinions, the dispute and its history, financial and remittance issues, legal and complaint actions, and the current status/outcomes. Use document names where helpful.",
 "study_guide": "Create a study guide: key concepts, definitions, and 5 review questions.",
 "report": "Write a structured report: Title, Executive Summary, Key Findings, Analysis, Recommendations.",
 "slides": "Create a slide deck. Output each slide as '# <Slide Title>' followed by 3-5 short bullet '- ' lines, and separate slides with a line containing only '---'. 8-12 slides.",
 "mindmap": "Create a mind map as a markdown nested bullet list (use 2-space indentation per level), one central topic with branches.",
 "flashcards": "Create 12 flashcards. Format each as:\n**Q:** question\n**A:** answer\n(blank line between cards).",
 "quiz": "Create a 10-question multiple-choice quiz. Format each as:\n**Q1.** question\n- A) ...\n- B) ...\n- C) ...\n- D) ...\n**Answer:** X",
 "infographic": "Design an infographic as markdown: a '# Title', then 6-9 lines each '- **<stat or keyword>** — short description'.",
 "datatable": "Extract a data table as a markdown table with a header row and 6-12 rows relevant to the material.",
}


def _corpus(nb_id, limit=30000, top_k=30):
    n = _db().execute("SELECT COUNT(*) c FROM sources WHERE notebook_id=?", (nb_id,)).fetchone()["c"]
    if n == 0:
        raise HTTPException(404, "No sources in this notebook yet.")
    q = ("Summarize the entire workers' compensation case: parties, claim numbers, injuries and medical "
         "opinions, dispute history, financial and remittance issues, legal actions, complaints, forensic "
         "findings, outcomes and chronology")
    hits = rag.retrieve(nb_id, q, top_k=top_k)
    body = "\n\n".join(f"[{i+1}] ({h['name']})\n{h['text'][:1400]}" for i, h in enumerate(hits))
    return f"This notebook has {n} sources. Key excerpts across the case:\n\n{body}"[:limit]


def _make_artifact(nb_id, kind):
    if kind not in STUDIO:
        raise HTTPException(404, f"unknown studio item: {kind}")
    title, _icon = STUDIO[kind]
    corpus = _corpus(nb_id)
    if kind == "audio":
        path = audio.render_audio(nb_id, title, corpus)
        return title, f"/api/notebooks/{nb_id}/media/audio-overview.mp3"
    if kind == "video":
        slides_md = _complete("Create a slide deck. Output each slide as '# <Slide Title>' then 3 short '- ' bullets, separated by a line '---'. 6-8 slides.\n\n" + corpus, 3000)
        titles = [l.lstrip("# ").strip() for l in slides_md.splitlines() if l.strip().startswith("#")]
        slides = [{"title": t} for t in titles] or [{"title": "Overview"}]
        audio_path = audio.render_audio(nb_id, title, corpus)
        video.make_slideshow(nb_id, slides, audio_path)
        return title, f"/api/notebooks/{nb_id}/media/video-overview.mp4"
    content = _complete(_PROMPTS.get(kind, "Summarize the material.") + "\n\nMaterial:\n" + corpus, 3000)
    if not content:
        raise HTTPException(500, "generation returned no content")
    return title, content


@app.post("/api/notebooks/{nb_id}/studio")
def studio(nb_id: str, body: StudioIn):
    title, content = _make_artifact(nb_id, body.kind)
    did = uuid.uuid4().hex
    c = _db()
    c.execute("INSERT INTO documents VALUES (?,?,?,?,?,datetime('now'))", (did, nb_id, body.kind, title, content))
    c.commit()
    return {"id": did, "kind": body.kind, "title": title, "icon": STUDIO[body.kind][1], "created": "now"}


@app.get("/api/notebooks/{nb_id}/documents")
def list_documents(nb_id: str):
    c = _db()
    rows = c.execute("SELECT id, kind, title, created FROM documents WHERE notebook_id=? ORDER BY created DESC", (nb_id,)).fetchall()
    n = c.execute("SELECT COUNT(*) n FROM sources WHERE notebook_id=?", (nb_id,)).fetchone()["n"]
    return [{**dict(r), "icon": STUDIO.get(r["kind"], ("", "description"))[1], "sources": n} for r in rows]


@app.get("/api/documents/{did}")
def get_document(did: str):
    r = _db().execute("SELECT * FROM documents WHERE id=?", (did,)).fetchone()
    if not r:
        raise HTTPException(404, "document not found")
    d = dict(r)
    d["icon"] = STUDIO.get(d["kind"], ("", "description"))[1]
    return d


@app.delete("/api/documents/{did}")
def del_document(did: str):
    c = _db()
    c.execute("DELETE FROM documents WHERE id=?", (did,))
    c.commit()
    return {"deleted": did}


@app.get("/api/notebooks/{nb_id}/media/{name}")
def media(nb_id: str, name: str):
    p = os.path.join(DATA_DIR, "audio", nb_id, name) if name.endswith(".mp3") else os.path.join(DATA_DIR, "video", nb_id, name)
    if not os.path.exists(p):
        raise HTTPException(404, "not generated")
    mt = "audio/mpeg" if name.endswith(".mp3") else "video/mp4"
    return FileResponse(p, media_type=mt)